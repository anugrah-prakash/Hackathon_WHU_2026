import copy
from io import BytesIO
import json
import unittest
from unittest.mock import patch
import urllib.error

from medeasy.explain import (ExplanationError, explain, parse_explanation,
                             render_text, validate_source, exact_number_lines)
from medeasy.documents import extract_document

SOURCE = "Discharge letter.\nReference: MED-2026-0042\nVisit your GP by 21 October 2026.\nPay £15.00 by 28 October 2026."


def claim(text, quote):
    return {"text": text, "evidence": [quote] if quote else []}


def response():
    return {
        "summary": claim("This is your discharge letter.", "Discharge letter."),
        "action_required": [claim("Visit your GP by 21 October 2026.", "Visit your GP by 21 October 2026.")],
        "important_details": [claim("Reference: MED-2026-0042", "Reference: MED-2026-0042")],
    }


class GroundingTests(unittest.TestCase):
    def test_valid_three_sections_and_exact_numeric_backstop(self):
        parsed = parse_explanation(json.dumps(response()), SOURCE)
        text = render_text(parsed, SOURCE)
        self.assertEqual([line for line in text.splitlines() if line[:2] in ("1.", "2.", "3.")],
                         ["1. SUMMARY", "2. ACTION REQUIRED", "3. IMPORTANT DETAILS"])
        self.assertIn("Pay £15.00 by 28 October 2026.", text)
        self.assertIn("MED-2026-0042", text)

    def test_rejects_invented_or_reformatted_number(self):
        for text in ("Visit on 22 October 2026.", "Visit on 21/10/2026."):
            with self.subTest(text=text):
                data = response()
                data["action_required"][0]["text"] = text
                with self.assertRaises(ExplanationError):
                    parse_explanation(json.dumps(data), SOURCE)

    def test_rejects_invented_quote(self):
        data = response()
        data["summary"]["evidence"] = ["A different letter."]
        with self.assertRaises(ExplanationError):
            parse_explanation(json.dumps(data), SOURCE)

    def test_missing_information_uses_exact_phrase(self):
        data = response()
        data["important_details"].append(claim("Medicine dose: Not specified in the letter", None))
        self.assertIn("Not specified in the letter", render_text(parse_explanation(json.dumps(data), SOURCE), SOURCE))

    def test_unquoted_claim_rejected(self):
        data = response()
        data["action_required"][0] = claim("No payment is needed.", None)
        with self.assertRaises(ExplanationError):
            parse_explanation(json.dumps(data), SOURCE)

    def test_extra_section_rejected(self):
        data = response()
        data["conclusion"] = "Extra"
        with self.assertRaises(ExplanationError):
            parse_explanation(json.dumps(data), SOURCE)

    def test_empty_section_rejected(self):
        data = response()
        data["action_required"] = []
        with self.assertRaises(ExplanationError):
            parse_explanation(json.dumps(data), SOURCE)

    def test_multisentence_summary_keeps_overflow_in_details(self):
        data = response()
        data["summary"] = claim("This is your discharge letter. Visit your GP by 21 October 2026.",
                                "Discharge letter.\nReference: MED-2026-0042\nVisit your GP by 21 October 2026.")
        result = parse_explanation(json.dumps(data), SOURCE)
        self.assertEqual(result.summary.text, "This is your discharge letter.")
        self.assertEqual(result.details[0].text, "Visit your GP by 21 October 2026.")
        self.assertEqual(result.details[0].evidence, result.summary.evidence)

    def test_html_or_newline_rejected(self):
        for text in ("<script>alert(1)</script>", "Hello\n# New heading"):
            data = response()
            data["summary"]["text"] = text
            with self.assertRaises(ExplanationError):
                parse_explanation(json.dumps(data), SOURCE)

    def test_not_json_rejected(self):
        with self.assertRaises(ExplanationError):
            parse_explanation("```json\n{}\n```", SOURCE)

    def test_equivalent_number_encodings_accepted(self):
        euro_source = "Invoice: pay €1,500 by 28 October 2026.\nTake 10–15 mg daily.\nFee 1.500,00 recorded."
        data = {
            "summary": claim("The letter is one invoice.", "Invoice: pay €1,500 by 28 October 2026."),
            "action_required": [
                claim("You must pay 1500 by 28 October 2026.", euro_source),
                claim("Dose range is 10-15 mg daily.", euro_source),
            ],
            "important_details": [claim("Fee is 1.500,00.", euro_source)],
        }
        result = parse_explanation(json.dumps(data), euro_source)
        self.assertIn("1500", result.actions[0].text)

    def test_currency_symbol_treated_as_encoding(self):
        # Deliberate policy: a differing inline currency symbol counts as a
        # formatting variant, because the numeric value itself must still be
        # copied exactly from the quoted passage.
        data = response()
        data["action_required"] = [claim("Pay $15.00.", "Pay £15.00 by 28 October 2026.")]
        parse_explanation(json.dumps(data), SOURCE)

    def test_decimal_comma_never_merges_digits(self):
        source = "Take 1,5 mg daily."
        data = response()
        data["action_required"] = [claim("Take 15 mg daily.", source)]
        with self.assertRaises(ExplanationError):
            parse_explanation(json.dumps(data), source)

    def test_money_preserved(self):
        data = response()
        data["action_required"] = [claim("Pay £15.00.", "Pay £15.00 by 28 October 2026.")]
        parse_explanation(json.dumps(data), SOURCE)
        data["action_required"][0]["text"] = "Pay £150.00."
        with self.assertRaises(ExplanationError):
            parse_explanation(json.dumps(data), SOURCE)

    def test_numeric_lines_preserve_order_and_spelling(self):
        self.assertEqual(exact_number_lines("No numbers\nRef: 00A-2\nCost: € 1.234,50\nRef: 00A-2"),
                         ("Ref: 00A-2", "Cost: € 1.234,50"))

    def test_blank_or_oversize_source_rejected(self):
        for source in ("  ", "a" * 40001):
            with self.assertRaises(ExplanationError):
                validate_source(source)


class APITests(unittest.TestCase):
    def fake_opener(self, body):
        class Opener:
            def open(inner, request, timeout):
                self.request = request
                self.timeout = timeout
                return BytesIO(json.dumps(body).encode())
        return Opener()

    def test_api_request_and_response(self):
        indexed = response()
        for item in [indexed["summary"], *indexed["action_required"], *indexed["important_details"]]:
            item["evidence"] = ["S0001"]
        body = {"choices": [{"finish_reason": "stop", "message": {"content": json.dumps(indexed)}}]}
        with patch("urllib.request.build_opener", return_value=self.fake_opener(body)):
            result = explain(SOURCE, "secret-key", "example/model")
        payload = json.loads(self.request.data)
        self.assertEqual(self.request.full_url, "https://api.trustedtokens.eu/v1/chat/completions")
        self.assertEqual(self.request.get_header("Authorization"), "Bearer secret-key")
        self.assertEqual(json.loads(payload["messages"][1]["content"])["source_passages"], [{"id": "S0001", "text": SOURCE}])
        self.assertEqual(payload["response_format"]["json_schema"]["schema"]["properties"]["summary"]["properties"]["evidence"]["items"]["enum"], ["S0001"])
        self.assertEqual(payload["response_format"]["type"], "json_schema")
        self.assertEqual(result.summary.text, "This is your discharge letter.")

    def test_truncated_response_never_displayed(self):
        body = {"choices": [{"finish_reason": "length", "message": {"content": json.dumps(response())}}]}
        with patch("urllib.request.build_opener", return_value=self.fake_opener(body)):
            with self.assertRaisesRegex(ExplanationError, "incomplete"):
                explain(SOURCE, "secret-key", "example/model")

    def test_provider_error_does_not_expose_key_or_body(self):
        error = urllib.error.HTTPError("https://example.invalid", 401, "secret patient data", {}, None)
        with patch("urllib.request.build_opener") as build:
            build.return_value.open.side_effect = error
            with self.assertRaises(ExplanationError) as caught:
                explain(SOURCE, "secret-key", "example/model")
        self.assertNotIn("secret", str(caught.exception))
        self.assertIn("key", str(caught.exception))

    def test_repair_note_carried_in_system_message(self):
        indexed = response()
        for item in [indexed["summary"], *indexed["action_required"], *indexed["important_details"]]:
            item["evidence"] = ["S0001"]
        body = {"choices": [{"finish_reason": "stop", "message": {"content": json.dumps(indexed)}}]}
        with patch("urllib.request.build_opener", return_value=self.fake_opener(body)):
            explain(SOURCE, "secret-key", "example/model", repair_note="COPY NUMBERS EXACTLY.")
        payload = json.loads(self.request.data)
        self.assertIn("COPY NUMBERS EXACTLY.", payload["messages"][0]["content"])
        # The user payload with the letter text is unchanged by the note.
        self.assertEqual(json.loads(payload["messages"][1]["content"])["source_passages"][0]["text"], SOURCE)

    def test_no_request_for_missing_key_or_oversize(self):
        with patch("urllib.request.build_opener") as build:
            for source, key in ((SOURCE, ""), (SOURCE, "your_placeholder"), ("x" * 40001, "key")):
                with self.assertRaises(ExplanationError):
                    explain(source, key, "example/model")
            build.assert_not_called()

    def test_timeout_is_actionable(self):
        with patch("urllib.request.build_opener") as build:
            build.return_value.open.side_effect = TimeoutError()
            with self.assertRaisesRegex(ExplanationError, "connection"):
                explain(SOURCE, "key", "model")


class DocumentTests(unittest.TestCase):
    def test_utf8_bom(self):
        self.assertEqual(extract_document(b"\xef\xbb\xbfHello", "letter.txt"), "Hello")

    def test_unsupported_empty_corrupt_inputs(self):
        for data, name in ((b"", "a.txt"), (b"test", "a.exe"), (b"\xff", "a.txt"), (b"bad zip", "a.docx"), (b"bad pdf", "a.pdf")):
            with self.subTest(name=name), self.assertRaises(ExplanationError):
                extract_document(data, name)

    def test_oversize_file(self):
        with self.assertRaisesRegex(ExplanationError, "10 MB"):
            extract_document(b"x" * (10 * 1024 * 1024 + 1), "a.txt")

    def test_blank_pdf_is_not_silently_ignored(self):
        from pypdf import PdfWriter
        pdf = PdfWriter()
        pdf.add_blank_page(width=612, height=792)
        stream = BytesIO()
        pdf.write(stream)
        with self.assertRaisesRegex(ExplanationError, "scanned"):
            extract_document(stream.getvalue(), "letter.pdf")

    def test_encrypted_pdf_rejected(self):
        from pypdf import PdfWriter
        pdf = PdfWriter()
        pdf.add_blank_page(width=612, height=792)
        pdf.encrypt("password")
        stream = BytesIO()
        pdf.write(stream)
        with self.assertRaisesRegex(ExplanationError, "password"):
            extract_document(stream.getvalue(), "letter.pdf")

    def test_docx_preserves_table_order_and_header(self):
        from docx import Document
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = "Reference: MED-0001"
        doc.add_paragraph("Before table")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Medicine A"
        table.cell(0, 1).text = "5 mg"
        doc.add_paragraph("After table")
        stream = BytesIO()
        doc.save(stream)
        text = extract_document(stream.getvalue(), "letter.docx")
        self.assertIn("Reference: MED-0001", text)
        self.assertLess(text.index("Before table"), text.index("5 mg"))
        self.assertLess(text.index("5 mg"), text.index("After table"))

    def test_text_pdf_extraction(self):
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, NameObject, DictionaryObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 50 700 Td (Reference: MED-2026-0042) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
        stream = BytesIO()
        writer.write(stream)
        self.assertIn("MED-2026-0042", extract_document(stream.getvalue(), "letter.pdf"))


if __name__ == "__main__":
    unittest.main()
